import csv
import json
from io import BytesIO, StringIO
from typing import Any

from openpyxl import load_workbook


def parse_csv(content: bytes) -> dict[str, Any]:
    text = content.decode("utf-8-sig")
    reader = csv.DictReader(StringIO(text))
    return {
        "kind": "table",
        "columns": reader.fieldnames or [],
        "rows": list(reader),
    }


def parse_json(content: bytes) -> dict[str, Any]:
    return {"kind": "json", "data": json.loads(content.decode("utf-8-sig"))}


def parse_xlsx(content: bytes) -> dict[str, Any]:
    workbook = load_workbook(BytesIO(content), data_only=False, read_only=True)
    sheets = [
        {"name": sheet.title, "rows": [list(row) for row in sheet.iter_rows(values_only=True)]}
        for sheet in workbook.worksheets
    ]
    return {"kind": "workbook", "sheets": sheets}

def parse_pptx(file_path: str) -> dict[str, Any]:
    from pptx import Presentation

    presentation = Presentation(file_path)
    blocks: list[str] = []
    content_list: list[dict[str, Any]] = []
    for number, slide in enumerate(presentation.slides, start=1):
        blocks.append(f"## Slide {number}")
        content_list.append({"type": "title", "text": f"Slide {number}", "locator": {"kind": "slide", "slide": number}})
        for shape_index, shape in enumerate(slide.shapes, start=1):
            text = shape.text.strip() if hasattr(shape, "text") else ""
            if not text:
                continue
            is_title = shape == slide.shapes.title
            blocks.append(f"### {text}" if is_title else text)
            content_list.append(
                {
                    "type": "title" if is_title else "text",
                    "text": text,
                    "locator": {
                        "kind": "slide_region",
                        "slide": number,
                        "shape": shape_index,
                        "bbox": [int(shape.left), int(shape.top), int(shape.left + shape.width), int(shape.top + shape.height)],
                        "coordinate_space": "pptx_emu",
                    },
                }
            )
    return {"blocks": blocks, "content_list": content_list, "slide_count": len(presentation.slides)}